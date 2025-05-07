#!/usr/bin/env python3

import requests
import os
import re
import json
import logging
from datetime import datetime
from collections import Counter
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Selector to choose the function to run
mode = "reconcile"

# Base directory for all file paths
base_directory = "/Users/phobrla/Library/CloudStorage/OneDrive-Personal/Documents/Bulgarian Language Learning"

# Paths and directories
concatenated_file_path = os.path.join(base_directory, "concatenated.json")
query_parts_of_speech_json_path = os.path.join(base_directory, "Query Parts of Speech.json")

# Generate the timestamp suffix
timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
xlsx_output_file = os.path.join(base_directory, f"reconciliation_results_{timestamp}.xlsx")
log_file_path = os.path.join(base_directory, f"debug_{timestamp}.log")

# Ensure the output directory exists
os.makedirs(base_directory, exist_ok=True)

# Setup logging
logging.basicConfig(
    filename=log_file_path,
    level=logging.DEBUG,
    format="%(asctime)s - %(levelname)s - %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S"
)
logging.info("Script started")
logging.info(f"Mode: {mode}")
logging.info(f"XLSX Output File: {xlsx_output_file}")
logging.info(f"Log File: {log_file_path}")

# Strings to be cut off during the re-search
cutoff_strings = [
    " се", " [в]", " си", " [с]", " за", " в", " (се)", " (се) [с]", " (се) да"
]

# Summary data for logging
summary_data = Counter()


def reconcile_entries():
    logging.info("Starting reconcile mode.")
    
    # Load required files
    if not os.path.exists(concatenated_file_path):
        logging.error(f"Concatenated JSON file not found at: {concatenated_file_path}")
        print(f"Concatenated JSON file not found at {concatenated_file_path}. Please run 'concatenate' mode first.")
        return

    if not os.path.exists(query_parts_of_speech_json_path):
        logging.error(f"Query Parts of Speech JSON file not found at: {query_parts_of_speech_json_path}")
        print(f"Query Parts of Speech JSON file not found at {query_parts_of_speech_json_path}. Please ensure it exists.")
        return

    # Load concatenated data
    try:
        with open(concatenated_file_path, 'r', encoding='utf-8') as file:
            concatenated_data = json.load(file)
            concatenated_queries = set(entry["query"] for entry in concatenated_data)
        logging.debug(f"Total queries in concatenated.json: {len(concatenated_data)}")
    except (IOError, UnicodeDecodeError) as e:
        logging.error(f"Error reading concatenated file '{concatenated_file_path}': {e}")
        return

    # Load the Query Parts of Speech JSON file
    bulgarian_to_pos = {}
    try:
        with open(query_parts_of_speech_json_path, 'r', encoding='utf-8') as file:
            parts_of_speech_data = json.load(file).get("JSONdata", [])
            for entry in parts_of_speech_data:
                bulgarian_word = entry.get("Bulgarian", "").strip()
                part_of_speech = entry.get("Part of Speech", "").strip()
                if bulgarian_word and part_of_speech:
                    bulgarian_to_pos[bulgarian_word] = part_of_speech
        logging.debug(f"Total mappings in Query Parts of Speech.json: {len(bulgarian_to_pos)}")
    except (IOError, UnicodeDecodeError) as e:
        logging.error(f"Error reading JSON file '{query_parts_of_speech_json_path}': {e}")
        return

    # First Pass: Initial processing of queries
    results = []
    unmatched_revised_queries = []  # Collect unmatched revised queries for debugging
    for entry in concatenated_data:
        query = entry["query"]
        original_data = entry.get("data", None)

        # Default values for fields
        revised_query = ""
        cutoff_type = ""
        cutoff_applied = ""
        original_found_in_concatenated = "No"
        original_found_in_pos = "No"
        revised_status = ""
        revised_result = ""
        revised_part_of_speech = ""
        revised_found_in_concatenated = ""
        revised_found_in_pos = ""

        # Part of Speech
        original_part_of_speech = bulgarian_to_pos.get(query, "Unknown")
        if original_part_of_speech == "Unknown":
            logging.warning(f"Part of Speech for query '{query}' is unknown.")
            summary_data[f"Unknown Part of Speech: {query}"] += 1

        # Handle Verbs with cutoff logic
        if original_part_of_speech == "Verb":
            for cutoff in cutoff_strings:
                if query.endswith(cutoff):
                    revised_query = query[: -len(cutoff)].strip()
                    cutoff_type = "Verb Cutoff"
                    cutoff_applied = cutoff
                    break

        # Handle Adverbs and Unclassified Words ending in "но"
        elif original_part_of_speech in ["Adverb", "Unclassified Word"]:
            if query.endswith("но"):
                revised_query = re.sub(r"но$", "ен", query)
                cutoff_type = "Adverb Modification"
                cutoff_applied = "но → ен"

        # Determine "Found in" values for original query
        original_found_in_concatenated = "Yes" if query in concatenated_queries else "No"
        original_found_in_pos = "Yes" if query in bulgarian_to_pos else "No"

        # Validate Revised Query
        if revised_query:
            revised_found_in_concatenated = "Yes" if revised_query in concatenated_queries else "No"
            revised_found_in_pos = "Yes" if revised_query in bulgarian_to_pos else "No"
            revised_part_of_speech = bulgarian_to_pos.get(revised_query, "Unknown")
            revised_status = "Success" if revised_found_in_concatenated == "Yes" else "Failure"
            revised_result = "Match Found" if revised_found_in_pos == "Yes" else "No Match"

            # Log unmatched revised queries
            if revised_found_in_concatenated == "No":
                unmatched_revised_queries.append(revised_query)
        else:
            # Ensure blanks for revised fields if revised query is blank
            revised_status = ""
            revised_result = ""
            revised_found_in_concatenated = ""
            revised_found_in_pos = ""

        results.append([
            query, original_part_of_speech, original_found_in_concatenated, original_found_in_pos,
            revised_query, revised_part_of_speech, cutoff_type, cutoff_applied, revised_status, revised_result,
            revised_found_in_concatenated, revised_found_in_pos
        ])

    # Log unmatched queries
    if unmatched_revised_queries:
        logging.warning(f"Unmatched Revised Queries: {len(unmatched_revised_queries)}")
        for qr in unmatched_revised_queries:
            logging.warning(f"Unmatched Revised Query: {qr}")

    # Generate XLSX file
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    # Write headers
    headers = [
        "Original Query", "Original Part of Speech", "Original Found in Concatenated", "Original Found in POS",
        "Revised Query", "Revised Part of Speech", "Cutoff Type", "Cutoff Applied",
        "Revised Status", "Revised Result", "Revised Found in Concatenated", "Revised Found in POS"
    ]
    sheet.append(headers)

    # Write data
    for row in results:
        sheet.append(row)

    # Auto-fit column widths
    for column in sheet.columns:
        max_length = max(len(str(cell.value)) for cell in column if cell.value)
        column_letter = column[0].column_letter
        sheet.column_dimensions[column_letter].width = max_length + 2

    # Apply table style
    table_range = f"A1:L{sheet.max_row}"
    table = Table(displayName="Table1", ref=table_range)
    style = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    table.tableStyleInfo = style
    sheet.add_table(table)

    # Save the workbook
    try:
        workbook.save(xlsx_output_file)
        logging.info(f"Reconciliation completed. Results saved to {xlsx_output_file}")
        print(f"Reconciliation completed. Results saved to {xlsx_output_file}")
    except IOError as e:
        logging.error(f"Error writing reconciliation results to '{xlsx_output_file}': {e}")

# Main workflow
if mode == "reconcile":
    reconcile_entries()
else:
    logging.error(f"Unknown mode: {mode}")
    print(f"Unknown mode: {mode}")
    print("Available modes: reconcile")