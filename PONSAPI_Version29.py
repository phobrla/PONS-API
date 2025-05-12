#!/usr/bin/env python3

import os
import json
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from pyxlsb import open_workbook

# Selector to choose the function to run
mode = "reconcile"  # Default mode, can be changed to any available mode

# Base directory for all file paths
base_directory = "/Users/phobrla/Library/CloudStorage/OneDrive-Personal/Documents/Bulgarian Language Learning"

# File paths
concatenated_file_path = os.path.join(base_directory, "concatenated.json")
flashcards_xlsb_path = os.path.join(base_directory, "Flashcards.xlsb")

# Generate the timestamp suffix
timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
xlsx_output_file = os.path.join(base_directory, f"reconciliation_results_{timestamp}.xlsx")

# Ensure the output directory exists
os.makedirs(base_directory, exist_ok=True)


def load_anki_table_from_xlsb(file_path):
    """
    Load the 'Anki' table from the Flashcards.xlsb file in read-only mode.
    This function can parse the file even if it is open in Microsoft Excel.
    """
    anki_data = []

    try:
        with open_workbook(file_path) as workbook:
            with workbook.get_sheet("Anki") as sheet:
                rows = sheet.rows()
                headers = [cell.v for cell in next(rows)]  # Extract headers from the first row
                header_map = {header: idx for idx, header in enumerate(headers)}

                for row in rows:
                    row_data = {header: row[header_map[header]].v if header_map[header] < len(row) else None for header in headers}
                    anki_data.append(row_data)

    except Exception as e:
        print(f"An error occurred while reading the 'Anki' table: {e}")
    
    return anki_data


def reconcile_entries():
    """
    Reconciles entries between Flashcards.xlsb and concatenated.json.
    Processes 'Bulgarian 1' and 'Bulgarian 2' independently.
    """
    if not os.path.exists(concatenated_file_path):
        print(f"Concatenated JSON file not found at {concatenated_file_path}. Please run 'fetch' mode first.")
        return

    if not os.path.exists(flashcards_xlsb_path):
        print(f"Flashcards.xlsb file not found at {flashcards_xlsb_path}. Please ensure it exists.")
        return

    try:
        # Load concatenated.json
        with open(concatenated_file_path, 'r', encoding='utf-8') as file:
            concatenated_data = json.load(file)

        # Load the "Anki" table from Flashcards.xlsb
        anki_data = load_anki_table_from_xlsb(flashcards_xlsb_path)

        # Prepare reconciliation results
        results = []

        for row in anki_data:
            note_id = row.get("Note ID", "").strip() if row.get("Note ID") else ""
            bulgarian_1 = row.get("Bulgarian 1", "").strip() if row.get("Bulgarian 1") else ""
            bulgarian_2 = row.get("Bulgarian 2", "").strip() if row.get("Bulgarian 2") else ""
            part_of_speech = row.get("Part of Speech", "").strip() if row.get("Part of Speech") else ""

            # Process Bulgarian 1
            bulgarian_1_status = process_bulgarian_field(bulgarian_1, concatenated_data)

            # Process Bulgarian 2
            bulgarian_2_status = process_bulgarian_field(bulgarian_2, concatenated_data)

            # Append results
            results.append([
                note_id, bulgarian_1, bulgarian_1_status, bulgarian_2, bulgarian_2_status, part_of_speech
            ])

        # Create the Excel file
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reconciliation Results"

        # Write headers
        headers = [
            "Note ID", "Bulgarian 1", "Bulgarian 1 Status", 
            "Bulgarian 2", "Bulgarian 2 Status", "Part of Speech"
        ]
        sheet.append(headers)

        # Write data
        for row in results:
            sheet.append(row)

        # Auto-fit column widths
        for column in sheet.columns:
            max_length = max(len(str(cell.value)) for cell in column if cell.value)
            column_letter = column[0].column_letter
            sheet.column_dimensions[column_letter].width = max_length + 2

        # Apply table style
        table_range = f"A1:F{sheet.max_row}"
        table = Table(displayName="Table1", ref=table_range)
        style = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False
        )
        table.tableStyleInfo = style
        sheet.add_table(table)

        # Save the workbook
        try:
            workbook.save(xlsx_output_file)
            print(f"Reconciliation completed. Results saved to {xlsx_output_file}")
        except IOError as e:
            print(f"Error writing reconciliation results to '{xlsx_output_file}': {e}")

    except Exception as e:
        print(f"An error occurred: {e}")


def process_bulgarian_field(bulgarian_field, concatenated_data):
    """
    Processes a single Bulgarian field (Bulgarian 1 or Bulgarian 2).
    Checks its presence and status in concatenated.json.
    """
    if not bulgarian_field:
        return "Missing"

    # Search for the query in concatenated.json
    concatenated_entry = next((item for item in concatenated_data if item.get("query") == bulgarian_field), None)

    if not concatenated_entry:
        return "Missing"

    # Check the contents of the data field
    data = concatenated_entry.get("data", {})
    if isinstance(data, dict):
        # Explicitly check for the "Received status code 204" error
        if data.get("error") == "Received status code 204":
            return "Not Found"
    elif isinstance(data, list):
        # Check if there are no hits
        hits = [item.get("hits", []) for item in data]
        if not any(hits):  # No hits at all
            return "No Headword"
        # Check for roms in hits
        if any("roms" in hit for sublist in hits for hit in sublist):
            return "Found"

    # If none of the above conditions match, assume "No Headword"
    return "No Headword"


# New modules for additional modes
def fetch_data():
    print("Fetching data... (Implementation pending)")


def createacronyms():
    print("Expanding acronyms... (Implementation pending)")


def validate_data():
    print("Validating data... (Implementation pending)")


def process_data():
    print("Processing data... (Implementation pending)")


def schematize_data():
    print("Schematizing data... (Implementation pending)")


# Main workflow with dynamic mode selection
available_modes = {
    "reconcile": reconcile_entries,
    "fetch": fetch_data,
    "createacronyms": createacronyms,
    "validate": validate_data,
    "process": process_data,
    "schematize": schematize_data,
}

if mode in available_modes:
    available_modes[mode]()
else:
    print(f"Unknown mode: {mode}")
    print(f"Available modes: {', '.join(available_modes.keys())}")