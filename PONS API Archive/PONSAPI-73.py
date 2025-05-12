#!/usr/bin/env python3

import os
import re
import json
from datetime import datetime
from collections import Counter
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from pyxlsb import open_workbook

# Selector to choose the function to run
mode = "reconcile"

# Base directory for all file paths
base_directory = "/Users/phobrla/Library/CloudStorage/OneDrive-Personal/Documents/Bulgarian Language Learning"

# File paths
concatenated_file_path = os.path.join(base_directory, "concatenated.json")
flashcards_xlsb_path = os.path.join(base_directory, "Flashcards.xlsb")
schematized_file_path = os.path.join(base_directory, "schematized.json")

# Generate the timestamp suffix
timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
xlsx_output_file = os.path.join(base_directory, f"reconciliation_results_{timestamp}.xlsx")

# Ensure the output directory exists
os.makedirs(base_directory, exist_ok=True)

# Strings to be cut off during the re-search
cutoff_strings = [
    " се", " [в]", " си", " [с]", " за", " в", " (се)", " (се) [с]", " (се) да"
]

# Summary data for logging
summary_data = Counter()


def load_anki_table_from_xlsb(file_path):
    """
    Load the 'Anki' table from the Flashcards.xlsb file in read-only mode.
    This function can parse the file even if it is open in Microsoft Excel.
    """
    anki_data = []

    try:
        with open_workbook(file_path) as workbook:
            with workbook.get_sheet("Anki") as sheet:
                rows = sheet.rows()
                headers = [cell.v for cell in next(rows)]  # Extract headers from the first row
                header_map = {header: idx for idx, header in enumerate(headers)}

                for row in rows:
                    row_data = {header: row[header_map[header]].v if header_map[header] < len(row) else None for header in headers}
                    anki_data.append(row_data)

    except Exception as e:
        print(f"An error occurred while reading the 'Anki' table: {e}")
    
    return anki_data


def fetch_data():
    """
    Function to fetch raw data and save it to concatenated.json.
    This function can fetch data from an API, database, or other sources.
    """
    try:
        # Example data to simulate fetching
        fetched_data = [
            {"query": "корморан", "data": {"response_text": "Some response", "error": None}},
            {"query": "слон", "data": {"response_text": "", "error": "Received status code 204"}},
            {"query": "тигър", "data": {"response_text": "Another response", "error": None}}
        ]

        # Save fetched data to concatenated.json
        with open(concatenated_file_path, 'w', encoding='utf-8') as outfile:
            json.dump(fetched_data, outfile, ensure_ascii=False, indent=4)

        print(f"Data fetching completed. Fetched data saved to {concatenated_file_path}")

    except Exception as e:
        print(f"An error occurred during data fetching: {e}")


def concatenate_files():
    """
    Function to concatenate multiple files into a single JSON file.
    This function is a placeholder and should be implemented if needed.
    """
    print("Concatenate functionality is not implemented yet.")
    # Implement concatenation logic here if needed


def reconcile_entries():
    """
    Function to reconcile entries between the concatenated.json file and the Anki flashcards table.
    This function also processes the concatenated.json file before reconciliation.
    """
    if not os.path.exists(concatenated_file_path):
        print(f"Concatenated JSON file not found at {concatenated_file_path}. Please run 'fetch' mode first.")
        return

    if not os.path.exists(flashcards_xlsb_path):
        print(f"Flashcards.xlsb file not found at {flashcards_xlsb_path}. Please ensure it exists.")
        return

    try:
        # Load and process concatenated data
        with open(concatenated_file_path, 'r', encoding='utf-8') as file:
            concatenated_data = json.load(file)

        processed_data = []
        for entry in concatenated_data:
            # Ensure each entry is a dictionary
            if isinstance(entry, dict):
                query = entry.get("query", "")
                data = entry.get("data", {})

                # Check if the entry is "Example Only"
                is_example_only = False
                if isinstance(data, list):
                    for item in data:
                        hits = item.get("hits", [])
                        if any('<span class="example">' in hit.get("source", "") for hit in hits):
                            is_example_only = True
                            break

                # Validate entry
                is_valid = isinstance(data, dict) and not data.get("error") and data.get("response_text", "").strip()
                processed_data.append({
                    "query": query,
                    "is_valid": is_valid,
                    "is_example_only": is_example_only
                })
            else:
                print(f"Unexpected entry format: {entry}")

        # Load the "Anki" table from Flashcards.xlsb
        anki_data = load_anki_table_from_xlsb(flashcards_xlsb_path)

        # Process both Bulgarian 1 and Bulgarian 2 as separate queries
        expanded_data = []
        for row in anki_data:
            note_id = row.get("Note ID", "").strip() if row.get("Note ID") else ""
            part_of_speech = row.get("Part of Speech", "").strip() if row.get("Part of Speech") else ""
            bulgarian_1 = row.get("Bulgarian 1", "").strip() if row.get("Bulgarian 1") else ""
            bulgarian_2 = row.get("Bulgarian 2", "").strip() if row.get("Bulgarian 2") else ""

            if bulgarian_1:
                expanded_data.append({
                    "query": bulgarian_1,
                    "note_id": note_id,
                    "part_of_speech": part_of_speech,
                })
            if bulgarian_2:
                expanded_data.append({
                    "query": bulgarian_2,
                    "note_id": note_id,
                    "part_of_speech": part_of_speech,
                })

        # Create mappings for queries
        query_to_pos = {entry["query"]: entry["part_of_speech"] for entry in expanded_data}
        query_to_note_id = {entry["query"]: entry["note_id"] for entry in expanded_data}

        # Process data
        results = []
        unmatched_revised_queries = []  # Collect unmatched revised queries for debugging

        for entry in processed_data:
            query = entry["query"]
            is_valid = entry["is_valid"]
            is_example_only = entry["is_example_only"]

            # Default values for fields
            revised_query = ""
            cutoff_type = ""
            cutoff_applied = ""
            original_found_in_concatenated = "Example Only" if is_example_only else ("Yes" if is_valid else "No")
            revised_found_in_concatenated = "Example Only" if is_example_only else ""
            original_found_in_pos = "No"
            revised_found_in_pos = ""
            revised_status = ""
            revised_result = ""
            revised_part_of_speech = ""
            original_note_id = ""
            revised_note_id = ""

            # Part of Speech
            original_part_of_speech = query_to_pos.get(query, "Unknown")
            original_note_id = query_to_note_id.get(query, "")

            # Handle Verbs with cutoff logic
            if original_part_of_speech == "Verb":
                for cutoff in cutoff_strings:
                    if query.endswith(cutoff):
                        revised_query = query[: -len(cutoff)].strip()
                        cutoff_type = "Verb Cutoff"
                        cutoff_applied = cutoff
                        break

            # Handle Adverbs and Unclassified Words
            elif original_part_of_speech in ["Adverb", "Unclassified Word"]:
                if query.endswith("ено"):
                    revised_query = re.sub(r"ено$", "ен", query)
                    cutoff_type = "Adverb Modification"
                    cutoff_applied = "ено → ен"
                elif query.endswith("но"):
                    revised_query = re.sub(r"но$", "ен", query)
                    cutoff_type = "Adverb Modification"
                    cutoff_applied = "но → ен"
                elif query.endswith("о"):
                    revised_query = re.sub(r"о$", "ен", query)
                    cutoff_type = "Adverb Modification"
                    cutoff_applied = "о → ен"

            # Determine "Found in POS" for revised query
            if revised_query:
                revised_found_in_concatenated = "No"
                if revised_query in query_to_pos:
                    revised_found_in_pos = "Yes"
            else:
                revised_found_in_concatenated = ""
                revised_found_in_pos = ""

            if query in query_to_pos:
                original_found_in_pos = "Yes"

            # Validate Revised Query
            if revised_query:
                revised_part_of_speech = query_to_pos.get(revised_query, "Unknown")
                revised_note_id = query_to_note_id.get(revised_query, "")
                revised_status = "Success" if revised_found_in_concatenated == "Yes" else "Failure"
                revised_result = "Match Found" if revised_found_in_pos == "Yes" else "No Match"

                # Log unmatched revised queries
                if revised_found_in_concatenated == "No":
                    unmatched_revised_queries.append(revised_query)
            else:
                # Ensure blanks for revised fields if revised query is blank
                revised_status = ""
                revised_result = ""

            results.append([
                original_note_id, query, original_part_of_speech, 
                revised_query, revised_part_of_speech, cutoff_type, cutoff_applied,
                revised_status, revised_result, 
                original_found_in_concatenated, revised_found_in_concatenated, 
                original_found_in_pos, revised_found_in_pos, 
                revised_note_id
            ])

        # Create the Excel file only after all processing is complete
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        # Write headers
        headers = [
            "Original Note ID", "Original Query", "Original Part of Speech", 
            "Revised Query", "Revised Part of Speech", "Cutoff Type", 
            "Cutoff Applied", "Revised Status", "Revised Result", 
            "Original Found in Concatenated", "Revised Found in Concatenated", 
            "Original Found in POS", "Revised Found in POS", 
            "Revised Note ID"
        ]
        sheet.append(headers)

        # Write data
        for row in results:
            sheet.append(row)

        # Auto-fit column widths
        for column in sheet.columns:
            max_length = max(len(str(cell.value)) for cell in column if cell.value)
            column_letter = column[0].column_letter
            sheet.column_dimensions[column_letter].width = max_length + 2

        # Apply table style
        table_range = f"A1:N{sheet.max_row}"
        table = Table(displayName="Table1", ref=table_range)
        style = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False
        )
        table.tableStyleInfo = style
        sheet.add_table(table)

        # Save the workbook
        try:
            workbook.save(xlsx_output_file)
            print(f"Reconciliation completed. Results saved to {xlsx_output_file}")
        except IOError as e:
            print(f"Error writing reconciliation results to '{xlsx_output_file}': {e}")

    except Exception as e:
        print(f"An error occurred: {e}")


def schematize_entries():
    """
    Function to schematize the concatenated.json data into a structured JSON schema.
    """
    if not os.path.exists(concatenated_file_path):
        print(f"Concatenated JSON file not found at {concatenated_file_path}. Please run 'fetch' mode first.")
        return

    try:
        with open(concatenated_file_path, 'r', encoding='utf-8') as file:
            concatenated_data = json.load(file)

        schematized_data = []
        for entry in concatenated_data:
            query = entry.get("query", "")
            data = entry.get("data", {})

            # Create a structured schema for each entry
            schematized_entry = {
                "query": query,
                "response_text": data.get("response_text", ""),
                "error": data.get("error", None)
            }
            schematized_data.append(schematized_entry)

        # Write schematized data to schematized.json
        with open(schematized_file_path, 'w', encoding='utf-8') as outfile:
            json.dump(schematized_data, outfile, ensure_ascii=False, indent=4)

        print(f"Schematization completed. Schematized data saved to {schematized_file_path}")

    except Exception as e:
        print(f"An error occurred during schematization: {e}")


# Main workflow
if mode == "fetch":
    fetch_data()
elif mode == "reconcile":
    reconcile_entries()
elif mode == "concatenate":
    concatenate_files()
elif mode == "schematize":
    schematize_entries()
else:
    print(f"Unknown mode: {mode}")
    print("Available modes: fetch, reconcile, concatenate, schematize")